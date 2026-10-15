"""Excel workbook generator for task-aware secure analyses.

Builds MEDHA_<task label>_Secure_Analysis.xlsx with 8 sheets from the
task-derived AnalysisResult payload.
"""

from __future__ import annotations

import io
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


HEADER_FILL = PatternFill(start_color="1a2332", end_color="1a2332", fill_type="solid")
HEADER_FONT = Font(name="Calibri", bold=True, color="60a5fa", size=11)
DATA_FONT = Font(name="Calibri", color="e5e7eb", size=10)
TITLE_FONT = Font(name="Calibri", bold=True, color="f3f4f6", size=14)
SUBTITLE_FONT = Font(name="Calibri", color="9ca3af", size=10)
THIN_BORDER = Border(
    bottom=Side(style="thin", color="374151"),
)


def _style_header(ws, row: int, cols: int) -> None:
    for col in range(1, cols + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="left", vertical="center")
        cell.border = THIN_BORDER


def _auto_width(ws, min_width: int = 10, max_width: int = 40) -> None:
    for col in ws.columns:
        max_len = min_width
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.value:
                max_len = max(max_len, min(len(str(cell.value)) + 2, max_width))
        ws.column_dimensions[col_letter].width = max_len


def generate_excel(analysis_data: dict) -> bytes:
    """Generate the 8-sheet Excel workbook for a task-aware analysis."""
    wb = Workbook()

    equipment = analysis_data.get("equipment_info", {})
    failure_summary = analysis_data.get("failure_summary", {})
    sensor_summary = analysis_data.get("sensor_summary", {})
    maintenance_summary = analysis_data.get("maintenance_summary", {})
    risk = analysis_data.get("risk_assessment", {})

    # --- Sheet 1: Summary ---
    ws = wb.active
    ws.title = "Summary"
    ws.sheet_properties.tabColor = "3b82f6"

    ws.cell(
        row=1, column=1,
        value=f"MEDHA — {analysis_data.get('title', 'Task')} Secure Analysis",
    ).font = TITLE_FONT
    ws.cell(
        row=2, column=1,
        value="Synthetic demonstration data — Not MRPL internal data",
    ).font = SUBTITLE_FONT
    ws.cell(
        row=3, column=1,
        value=f"Generated: {datetime.utcnow().isoformat()}",
    ).font = SUBTITLE_FONT

    findings = failure_summary.get("findings", [])
    rows = [
        ("Task ID", analysis_data.get("task_id", "")),
        ("Analysis ID", analysis_data.get("analysis_id", "")),
        ("Security Mode", analysis_data.get("security_mode", equipment.get("security_mode", ""))),
        ("Task Type", analysis_data.get("task_type", "")),
        ("Equipment Context ID", equipment.get("id", "")),
        ("Status", equipment.get("status", "")),
        ("Sensitivity", equipment.get("sensitivity", "")),
        ("Network Policy", equipment.get("network_policy", "")),
        ("Findings", len(findings)),
        ("Blocked Actions", failure_summary.get("blocked_actions", 0)),
        ("Messages", sensor_summary.get("messages_count", 0)),
        ("Attachments", sensor_summary.get("attachments_count", 0)),
        ("Model Used", sensor_summary.get("model_used", "")),
        ("Risk Level", risk.get("risk_level", "")),
        ("Risk Score", risk.get("risk_score", 0)),
    ]

    ws.cell(row=5, column=1, value="Metric").font = HEADER_FONT
    ws.cell(row=5, column=2, value="Value").font = HEADER_FONT
    _style_header(ws, 5, 2)

    for i, (label, value) in enumerate(rows, start=6):
        ws.cell(row=i, column=1, value=label).font = DATA_FONT
        ws.cell(row=i, column=2, value=value).font = DATA_FONT
        ws.cell(row=i, column=1).border = THIN_BORDER
        ws.cell(row=i, column=2).border = THIN_BORDER

    _auto_width(ws)

    # --- Sheet 2: Findings ---
    ws2 = wb.create_sheet("Findings")
    ws2.sheet_properties.tabColor = "ef4444"
    headers = ["Category", "Severity", "Label", "Detail", "Source"]
    for col, h in enumerate(headers, 1):
        ws2.cell(row=1, column=col, value=h)
    _style_header(ws2, 1, len(headers))

    for i, f in enumerate(findings, start=2):
        ws2.cell(row=i, column=1, value=f.get("category", "")).font = DATA_FONT
        ws2.cell(row=i, column=2, value=f.get("severity", "")).font = DATA_FONT
        ws2.cell(row=i, column=3, value=f.get("label", "")).font = DATA_FONT
        ws2.cell(row=i, column=4, value=f.get("detail", "")).font = DATA_FONT
        ws2.cell(row=i, column=5, value=f.get("source", "")).font = DATA_FONT
        for col in range(1, len(headers) + 1):
            ws2.cell(row=i, column=col).border = THIN_BORDER
    if not findings:
        ws2.cell(row=2, column=1, value="No findings recorded").font = DATA_FONT

    _auto_width(ws2)

    # --- Sheet 3: Task Context ---
    ws3 = wb.create_sheet("Task Context")
    ws3.sheet_properties.tabColor = "22c55e"
    context_rows = [
        ("Messages", sensor_summary.get("messages_count", 0)),
        ("Attachments", sensor_summary.get("attachments_count", 0)),
        ("Authorized Sources", sensor_summary.get("sources_authorized", 0)),
        ("Authorized Tools", sensor_summary.get("tools_authorized", 0)),
        ("Blocked Tools", sensor_summary.get("tools_blocked", 0)),
        ("Model Used", sensor_summary.get("model_used", "")),
        ("Purpose", equipment.get("purpose", "")),
    ]
    ws3.cell(row=1, column=1, value="Task Context").font = TITLE_FONT
    ws3.cell(row=3, column=1, value="Metric").font = HEADER_FONT
    ws3.cell(row=3, column=2, value="Value").font = HEADER_FONT
    _style_header(ws3, 3, 2)
    for i, (label, value) in enumerate(context_rows, start=4):
        ws3.cell(row=i, column=1, value=label).font = DATA_FONT
        ws3.cell(row=i, column=2, value=value).font = DATA_FONT
        ws3.cell(row=i, column=1).border = THIN_BORDER
        ws3.cell(row=i, column=2).border = THIN_BORDER
    _auto_width(ws3)

    # --- Sheet 4: Execution Steps ---
    ws4 = wb.create_sheet("Execution Steps")
    ws4.sheet_properties.tabColor = "f59e0b"
    step_headers = ["Tool", "Action", "Status"]
    for col, h in enumerate(step_headers, 1):
        ws4.cell(row=1, column=col, value=h)
    _style_header(ws4, 1, len(step_headers))

    steps = maintenance_summary.get("executed_steps", [])
    for i, s in enumerate(steps, start=2):
        ws4.cell(row=i, column=1, value=s.get("tool", "")).font = DATA_FONT
        ws4.cell(row=i, column=2, value=s.get("action", "")).font = DATA_FONT
        ws4.cell(row=i, column=3, value=str(s.get("status", ""))).font = DATA_FONT
        for col in range(1, len(step_headers) + 1):
            ws4.cell(row=i, column=col).border = THIN_BORDER
    if not steps:
        ws4.cell(row=2, column=1, value="No authorized tools executed").font = DATA_FONT

    _auto_width(ws4)

    # --- Sheet 5: Risk Assessment ---
    ws5 = wb.create_sheet("Risk Assessment")
    ws5.sheet_properties.tabColor = "f97316"
    ws5.cell(row=1, column=1, value="Risk Assessment").font = TITLE_FONT
    risk_rows = [
        ("Risk Level", risk.get("risk_level", "")),
        ("Risk Score", risk.get("risk_score", 0)),
        ("Critical Findings", risk.get("critical_failures_count", 0)),
        ("High Findings", risk.get("high_failures_count", 0)),
        ("MTBF Trend", risk.get("mtbf_trend", "")),
        ("Recommendation Priority", risk.get("recommendation_priority", "")),
    ]
    ws5.cell(row=3, column=1, value="Metric").font = HEADER_FONT
    ws5.cell(row=3, column=2, value="Value").font = HEADER_FONT
    _style_header(ws5, 3, 2)
    for i, (label, value) in enumerate(risk_rows, start=4):
        ws5.cell(row=i, column=1, value=label).font = DATA_FONT
        ws5.cell(row=i, column=2, value=value).font = DATA_FONT
        ws5.cell(row=i, column=1).border = THIN_BORDER
        ws5.cell(row=i, column=2).border = THIN_BORDER
    _auto_width(ws5)

    # --- Sheet 6: Recommendations ---
    ws6 = wb.create_sheet("Recommendations")
    ws6.sheet_properties.tabColor = "8b5cf6"
    rec_headers = ["ID", "Priority", "Action", "Rationale", "Est. Cost (USD)", "Est. Downtime (h)", "Tools Required"]
    for col, h in enumerate(rec_headers, 1):
        ws6.cell(row=1, column=col, value=h)
    _style_header(ws6, 1, len(rec_headers))

    recommendations = analysis_data.get("recommendations", [])
    for i, rec in enumerate(recommendations, start=2):
        ws6.cell(row=i, column=1, value=rec.get("id", "")).font = DATA_FONT
        ws6.cell(row=i, column=2, value=rec.get("priority", "")).font = DATA_FONT
        ws6.cell(row=i, column=3, value=rec.get("action", "")).font = DATA_FONT
        ws6.cell(row=i, column=4, value=rec.get("rationale", "")).font = DATA_FONT
        ws6.cell(row=i, column=5, value=rec.get("estimated_cost_usd", 0)).font = DATA_FONT
        ws6.cell(row=i, column=6, value=rec.get("estimated_downtime_hours", 0)).font = DATA_FONT
        ws6.cell(row=i, column=7, value=", ".join(rec.get("tools_required", []))).font = DATA_FONT
        for col in range(1, len(rec_headers) + 1):
            ws6.cell(row=i, column=col).border = THIN_BORDER
    _auto_width(ws6)

    # --- Sheet 7: Graph Data ---
    ws7 = wb.create_sheet("Graph Data")
    ws7.sheet_properties.tabColor = "06b6d4"
    graph_data = analysis_data.get("graph_data", {})

    row = 1
    for chart_name, chart_data in graph_data.items():
        ws7.cell(row=row, column=1, value=chart_data.get("title", chart_name)).font = TITLE_FONT
        row += 1
        ws7.cell(row=row, column=1, value="Label").font = HEADER_FONT
        ws7.cell(row=row, column=2, value="Value").font = HEADER_FONT
        _style_header(ws7, row, 2)
        row += 1
        labels = chart_data.get("labels", [])
        values = chart_data.get("values", [])
        for label, val in zip(labels, values):
            ws7.cell(row=row, column=1, value=label).font = DATA_FONT
            ws7.cell(row=row, column=2, value=val).font = DATA_FONT
            ws7.cell(row=row, column=1).border = THIN_BORDER
            ws7.cell(row=row, column=2).border = THIN_BORDER
            row += 1
        row += 2

    _auto_width(ws7)

    # --- Sheet 8: Security Pipeline ---
    ws8 = wb.create_sheet("Security Pipeline")
    ws8.sheet_properties.tabColor = "10b981"
    ws8.cell(row=1, column=1, value="Security Pipeline Trace").font = TITLE_FONT
    ws8.cell(
        row=2, column=1,
        value="All governance stages executed — blocked actions demonstrated",
    ).font = SUBTITLE_FONT

    pipe_headers = ["#", "Stage", "Status", "Detail", "Actor", "Latency (ms)", "Timestamp"]
    for col, h in enumerate(pipe_headers, 1):
        ws8.cell(row=4, column=col, value=h)
    _style_header(ws8, 4, len(pipe_headers))

    events = analysis_data.get("pipeline_events", [])
    for i, evt in enumerate(events, start=5):
        ws8.cell(row=i, column=1, value=i - 4).font = DATA_FONT
        ws8.cell(row=i, column=2, value=evt.get("stage", "").replace("_", " ")).font = DATA_FONT
        ws8.cell(row=i, column=3, value=evt.get("status", "").upper()).font = DATA_FONT
        ws8.cell(row=i, column=4, value=evt.get("detail", "")).font = DATA_FONT
        ws8.cell(row=i, column=5, value=evt.get("actor", "")).font = DATA_FONT
        ws8.cell(row=i, column=6, value=evt.get("latency_ms", 0)).font = DATA_FONT
        ws8.cell(row=i, column=7, value=evt.get("timestamp", "")).font = DATA_FONT
        for col in range(1, len(pipe_headers) + 1):
            ws8.cell(row=i, column=col).border = THIN_BORDER

    _auto_width(ws8, min_width=12, max_width=60)

    # Write to buffer
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()