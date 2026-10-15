"""P8 focused tests — REAL DOCX report generation from MEDHA task/analysis/evidence.

Covers:
  1. DOCX generation succeeds (real python-docx, not renamed HTML)
  2. Valid DOCX ZIP structure + document.xml
  3. Task information present
  4. Analysis/findings present
  5. Evidence/source information present
  6. No fabricated source
  7. Unauthorized task access blocked
  8. Unauthorized evidence not included
  9. No secrets/internal paths in document
 10. Download endpoint correct content-type/filename
 11. Empty/missing analysis handled safely
 12-14. Existing behavior remains via regression files (separate runs)

MOCKED vs REAL: LLM transport mocked by conftest; DOCX bytes are real.
"""

from __future__ import annotations

import io
import zipfile

import pytest
from docx import Document
from fastapi.testclient import TestClient

DOCX_CT = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def _authorize(client: TestClient, task_id: str, headers: dict[str, str]) -> None:
    client.patch(
        f"/api/tasks/{task_id}/status",
        json={"new_status": "authorized", "detail": "P8 authorization"},
        headers=headers,
    )


def _create_task(
    client: TestClient,
    headers: dict[str, str],
    *,
    title: str = "Pump P-204 Inspection Analysis",
    allowed_tools: list[str] | None = None,
    allowed_sources: list[str] | None = None,
) -> str:
    payload = {
        "title": title,
        "purpose": "Evaluate Pump P-204 synthetic inspection evidence",
        "allowed_tools": allowed_tools
        if allowed_tools is not None
        else ["document_reader", "data_analyzer"],
        "allowed_sources": allowed_sources or [],
    }
    resp = client.post("/api/tasks", json=payload, headers=headers)
    assert resp.status_code in (200, 201), resp.text
    task_id = resp.json()["task_id"]
    _authorize(client, task_id, headers)
    return task_id


def _assert_real_docx(content: bytes) -> tuple[Document, str]:
    assert len(content) > 500, "DOCX too small"
    # ZIP / OOXML structure
    with zipfile.ZipFile(io.BytesIO(content)) as zf:
        names = set(zf.namelist())
        assert "[Content_Types].xml" in names
        assert "word/document.xml" in names
        bad = zf.testzip()
        assert bad is None, f"corrupt zip entry {bad}"
    doc = Document(io.BytesIO(content))
    text = "\n".join(p.text for p in doc.paragraphs)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                text += "\n" + cell.text
    return doc, text


class TestDocxGeneration:
    def test_generate_returns_valid_docx(
        self, client: TestClient, auth_headers: dict[str, str]
    ) -> None:
        task_id = _create_task(client, auth_headers)
        resp = client.post(
            "/api/analysis/docx",
            json={"task_id": task_id},
            headers=auth_headers,
        )
        assert resp.status_code == 200, resp.text
        assert resp.headers["content-type"].startswith(DOCX_CT)
        disposition = resp.headers.get("content-disposition", "")
        assert "attachment" in disposition
        assert disposition.endswith('.docx"') or ".docx" in disposition
        doc, text = _assert_real_docx(resp.content)
        assert "MEDHA" in text
        assert "Sovereign Industrial AI Workbench" in text
        assert "Pump P-204 Inspection Analysis" in text

    def test_task_information_present(
        self, client: TestClient, auth_headers: dict[str, str]
    ) -> None:
        task_id = _create_task(client, auth_headers)
        resp = client.post(
            "/api/analysis/docx", json={"task_id": task_id}, headers=auth_headers
        )
        _, text = _assert_real_docx(resp.content)
        assert task_id in text
        assert "Security mode" in text
        assert "Sensitivity" in text
        assert "1. Executive Summary" in text
        assert "2. Task Information" in text
        assert "7. Security / Provenance Note" in text

    def test_analysis_findings_present(
        self, client: TestClient, auth_headers: dict[str, str]
    ) -> None:
        task_id = _create_task(client, auth_headers)
        # run analysis first so findings/recommendations exist
        run = client.post(
            "/api/analysis/run", json={"task_id": task_id}, headers=auth_headers
        )
        assert run.status_code == 200, run.text
        resp = client.post(
            "/api/analysis/docx", json={"task_id": task_id}, headers=auth_headers
        )
        doc, text = _assert_real_docx(resp.content)
        assert "3. Analysis / Findings" in text
        assert "5. Analysis Result / Recommendations" in text
        # real analysis always produces a narrative summary
        assert "Secure analysis executed for task" in text
        assert len(doc.tables) >= 1  # findings table

    def test_evidence_included_when_authorized(
        self, client: TestClient, auth_headers: dict[str, str]
    ) -> None:
        from app.services.rag.retrieval import retrieval_service

        source = "Pump_P204_Synthetic_Inspection_Note.txt"
        retrieval_service.ingest_document(
            "Synthetic Pump P-204 note: elevated vibration detected during inspection.",
            source,
            classification="internal",
        )
        task_id = _create_task(client, auth_headers, allowed_sources=[source])
        resp = client.post(
            "/api/analysis/docx", json={"task_id": task_id}, headers=auth_headers
        )
        assert resp.status_code == 200, resp.text
        _, text = _assert_real_docx(resp.content)
        assert "4. Evidence" in text
        assert source in text
        assert "elevated vibration detected" in text

    def test_no_unauthorized_evidence(
        self, client: TestClient, auth_headers: dict[str, str]
    ) -> None:
        from app.services.rag.retrieval import retrieval_service

        allowed = "Authorized_Pump_Source.txt"
        forbidden = "OTHER_USER_Secret_Document.txt"
        retrieval_service.ingest_document(
            "Authorized pump note for P-204.", allowed, classification="internal"
        )
        retrieval_service.ingest_document(
            "Forbidden confidential other-user content.", forbidden, classification="internal"
        )
        task_id = _create_task(client, auth_headers, allowed_sources=[allowed])
        resp = client.post(
            "/api/analysis/docx", json={"task_id": task_id}, headers=auth_headers
        )
        _, text = _assert_real_docx(resp.content)
        assert allowed in text
        assert forbidden not in text
        assert "Forbidden confidential" not in text

    def test_no_secrets_or_internal_paths(
        self, client: TestClient, auth_headers: dict[str, str]
    ) -> None:
        task_id = _create_task(client, auth_headers)
        resp = client.post(
            "/api/analysis/docx", json={"task_id": task_id}, headers=auth_headers
        )
        _, text = _assert_real_docx(resp.content)
        lowered = text.lower()
        assert "medha_jwt_secret" not in lowered
        assert "password" not in lowered
        assert "bearer " not in lowered
        assert "database_url" not in lowered
        # no raw host filesystem paths
        assert "oneDrive" not in text
        assert "C:\\Users" not in text

    def test_unauthenticated_blocked(self, client: TestClient) -> None:
        resp = client.post("/api/analysis/docx", json={"task_id": "x"})
        assert resp.status_code == 401

    def test_cross_user_blocked(
        self,
        client: TestClient,
        auth_headers: dict[str, str],
        engineer_headers: dict[str, str],
    ) -> None:
        task_id = _create_task(client, auth_headers)
        resp = client.post(
            "/api/analysis/docx",
            json={"task_id": task_id},
            headers=engineer_headers,
        )
        assert resp.status_code == 403

    def test_missing_task_404(
        self, client: TestClient, auth_headers: dict[str, str]
    ) -> None:
        resp = client.post(
            "/api/analysis/docx",
            json={"task_id": "does-not-exist"},
            headers=auth_headers,
        )
        assert resp.status_code == 404

    def test_empty_analysis_handled_safely(
        self, client: TestClient, auth_headers: dict[str, str]
    ) -> None:
        # Force a clean task with no prior analysis; generation runs analysis
        # inline but evidence may be empty — document must still be valid and
        # must not fabricate sources.
        task_id = _create_task(client, auth_headers, allowed_tools=["document_reader"])
        resp = client.post(
            "/api/analysis/docx", json={"task_id": task_id}, headers=auth_headers
        )
        assert resp.status_code == 200
        _, text = _assert_real_docx(resp.content)
        # Either real analysis ran, or an explicit empty-state sentence appears
        assert "No authorized evidence chunks" in text or "Secure analysis executed" in text
        assert "Fabricated_Source" not in text

    def test_no_fabricated_sources_when_no_rag(
        self, client: TestClient, auth_headers: dict[str, str]
    ) -> None:
        task_id = _create_task(client, auth_headers, allowed_sources=[])
        resp = client.post(
            "/api/analysis/docx", json={"task_id": task_id}, headers=auth_headers
        )
        _, text = _assert_real_docx(resp.content)
        # empty allowlist → fail-closed → no evidence section sources
        assert "No authorized evidence chunks were available" in text

    def test_audit_event_recorded(
        self, client: TestClient, auth_headers: dict[str, str]
    ) -> None:
        from app.services.audit_service import audit_service

        audit_service.clear()
        task_id = _create_task(client, auth_headers)
        client.post(
            "/api/analysis/docx", json={"task_id": task_id}, headers=auth_headers
        )
        events = audit_service.get_events_for_task(task_id)
        types = {e.event_type for e in events}
        assert "docx_report_generated" in types
        gen = [e for e in events if e.event_type == "docx_report_generated"][-1]
        assert "bytes" in gen.detail

    def test_generator_rejects_nothing_but_returns_bytes_for_minimal_payload(self) -> None:
        from app.services.analysis.docx_generator import generate_docx

        data = generate_docx({"task_id": "t-min", "title": "Minimal"})
        assert data[:2] == b"PK"  # ZIP magic — real DOCX container
        _, text = _assert_real_docx(data)
        assert "Minimal" in text
        assert "No analysis findings recorded" in text
